"""List-of-lists table helpers and file routing."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable, TypeGuard
import csv
import json

from .common import R_CODEPAGE, W_CODEPAGE, UTF8_BOM
from ._mapping import MappingError, apply_cvt_dict_to_llist
from ._sequences import is_empty_list
from ._validation import is_empty

_MISSING = object()


class TableFormatError(ValueError):
    """Raised for unsupported or malformed table data."""


def add_ext(path: str | Path, ext: str) -> str:
    target = Path(path)
    normalized = ext if ext.startswith(".") else f".{ext}"
    if target.suffix.lower() == normalized.lower():
        return str(target.with_suffix(normalized))
    return str(target.with_name(target.name + normalized))


def is_llist(value: object) -> TypeGuard[list[list[Any]]]:
    return isinstance(value, list) and all(isinstance(row, list) for row in value)


def is_empty_llist(rows: object) -> bool:
    if not is_llist(rows):
        raise TypeError("llist must be list[list]")
    return all(is_empty_list(row) for row in rows)


def is_completed_llist(rows: object) -> bool:
    return is_llist(rows) and bool(rows) and not all(is_empty_list(row) for row in rows)


def trim_llist_sublists(
    rows: list[list[Any]],
    *,
    eol: Any = _MISSING,
    remove_empty: bool = True,
    trim_empty_tails: bool = True,
    preserve_first: bool = False,
) -> None:
    has_eol = eol is not _MISSING

    def trim_empty_tail(row: list[Any]) -> None:
        index = len(row) - 1
        while index >= 0 and is_empty(row[index]):
            index -= 1
        del row[index + 1 :]

    write_index = 0
    for row_index, row in enumerate(rows):
        if has_eol:
            try:
                eol_index = row.index(eol)
            except ValueError:
                pass
            else:
                del row[eol_index:]

        if trim_empty_tails:
            trim_empty_tail(row)

        if not remove_empty:
            continue

        if preserve_first and row_index == 0:
            rows[write_index] = row
            write_index += 1
            continue

        if not is_empty_list(row):
            rows[write_index] = row
            write_index += 1

    if remove_empty:
        del rows[write_index:]


def deduplicated_llist(
    rows: list[list[Any]],
    *,
    start: int | None = None,
    stop: int | None = None,
) -> list[list[Any]]:
    slicer = slice(start, stop)
    seen_hashable: set[tuple[Any, ...]] = set()
    seen_unhashable: list[tuple[Any, ...]] = []
    result: list[list[Any]] = []

    for row in rows:
        key = tuple(row[slicer])
        try:
            if key in seen_hashable:
                continue
            seen_hashable.add(key)
        except TypeError:
            if any(key == old_key for old_key in seen_unhashable):
                continue
            seen_unhashable.append(key)
        result.append(row)
    return result


def sorted_llist(
    rows: list[list[Any]],
    *,
    start: int | None = None,
    stop: int | None = None,
    sort_cvt: Callable[[Any], Any] | str | None = "auto",
    headtail: int = 0,
    eol: Any = _MISSING,
) -> list[list[Any]]:
    if not isinstance(headtail, int) or isinstance(headtail, bool):
        raise TypeError(f"headtail must be int. Got: {type(headtail).__name__}={headtail!r}")

    auto = False
    if sort_cvt is None:
        pass
    elif isinstance(sort_cvt, str):
        if sort_cvt == "auto":
            auto = True
            sort_cvt = None
        else:
            raise ValueError(f'sort_cvt must be callable, None, or "auto". Got: {sort_cvt!r}')
    elif not callable(sort_cvt):
        raise TypeError(f'sort_cvt must be callable, None, or "auto". Got: {type(sort_cvt).__name__}')

    has_eol = eol is not _MISSING
    full_slice = start is None and stop is None
    slice_obj = slice(start, stop)

    def make_key(cvt: Callable[[Any], Any] | None) -> Callable[[list[Any]], tuple[Any, ...]]:
        def cut(row: list[Any]) -> list[Any]:
            if not has_eol:
                return row if full_slice else row[start:stop]
            n = len(row)
            a, b, _ = slice_obj.indices(n)
            try:
                b = row.index(eol, a, b)
            except ValueError:
                pass
            return row if a == 0 and b == n else row[a:b]

        if cvt is None:
            return lambda row: tuple(cut(row))
        return lambda row: tuple(map(cvt, cut(row)))

    def sort_with_key(key: Callable[[list[Any]], tuple[Any, ...]]) -> list[list[Any]]:
        if headtail == 0:
            return sorted(rows, key=key)
        if headtail > 0:
            body = rows[headtail:]
            body.sort(key=key)
            return rows[:headtail] + body
        body = rows[:headtail]
        body.sort(key=key)
        return body + rows[headtail:]

    if auto:
        try:
            return sort_with_key(make_key(int))
        except (ValueError, TypeError):
            return sort_with_key(make_key(str))
    return sort_with_key(make_key(sort_cvt))


def _prepare_external_rows(
    rows: Any,
    *,
    source_name: str,
    none: bool | None = None,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    if not is_llist(rows):
        raise TableFormatError(f"{source_name} table must be list[list]")
    result = rows
    if eol is not None:
        trim_llist_sublists(
            result,
            eol=eol,
            remove_empty=skip_empty,
            trim_empty_tails=trim_empty_tails,
            preserve_first=preserve_first,
        )
    elif trim_empty_tails or skip_empty:
        trim_llist_sublists(
            result,
            remove_empty=skip_empty,
            trim_empty_tails=trim_empty_tails,
            preserve_first=preserve_first,
        )
    if none:
        result = [[cell if cell != "" else None for cell in row] for row in result]
    return result


def _apply_table_conversion(
    rows: list[list[Any]],
    cvt_dict: dict[Any, Any] | None,
    *,
    header: bool,
    source_name: str,
) -> None:
    if cvt_dict is None:
        return
    try:
        apply_cvt_dict_to_llist(rows, cvt_dict, header=header)
    except MappingError as err:
        raise TableFormatError(f"{source_name} conversion failed: {err}") from err


def read_llist_from_csv_file(
    filename: str | Path,
    cvt_dict: dict[Any, Any] | None = None,
    *,
    header: bool = False,
    delimiter: str = ",",
    codepage: str = R_CODEPAGE,
    none: bool = True,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    with open(filename, encoding=codepage, newline="") as file:
        rows = [list(row) for row in csv.reader(file, delimiter=delimiter)]
    rows = _prepare_external_rows(
        rows,
        source_name="CSV",
        none=none,
        skip_empty=skip_empty,
        preserve_first=preserve_first,
        eol=eol,
        trim_empty_tails=trim_empty_tails,
    )
    _apply_table_conversion(rows, cvt_dict, header=header, source_name="CSV")
    return rows


def read_llist_from_json_file(
    filename: str | Path,
    cvt_dict: dict[Any, Any] | None = None,
    *,
    header: bool = False,
    codepage: str = R_CODEPAGE,
    none: bool | None = None,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    with open(filename, encoding=codepage) as file:
        rows = json.load(file)
    rows = _prepare_external_rows(
        rows,
        source_name="JSON",
        none=none,
        skip_empty=skip_empty,
        preserve_first=preserve_first,
        eol=eol,
        trim_empty_tails=trim_empty_tails,
    )
    _apply_table_conversion(rows, cvt_dict, header=header, source_name="JSON")
    return rows


def read_llist_from_xml_file(
    filename: str | Path,
    cvt_dict: dict[Any, Any] | None = None,
    *,
    header: bool = False,
    codepage: str = R_CODEPAGE,
    none: bool | None = None,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    try:
        import xmltodict
    except ImportError as err:
        raise TableFormatError("xmltodict package is required for XML parsing") from err
    with open(filename, encoding=codepage) as file:
        parsed: Any = xmltodict.parse(file.read())
    rows = _prepare_external_rows(
        parsed,
        source_name="XML",
        none=none,
        skip_empty=skip_empty,
        preserve_first=preserve_first,
        eol=eol,
        trim_empty_tails=trim_empty_tails,
    )
    _apply_table_conversion(rows, cvt_dict, header=header, source_name="XML")
    return rows


def _normalize_spreadsheet_rows(
    rows: list[list[Any]],
    *,
    none: bool = True,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
    empty_string_to_none: bool = False,
) -> list[list[Any]]:
    rows = _prepare_external_rows(
        rows,
        source_name="spreadsheet",
        none=False,
        skip_empty=skip_empty,
        preserve_first=preserve_first,
        eol=eol,
        trim_empty_tails=trim_empty_tails,
    )
    if empty_string_to_none:
        if none:
            rows = [[cell if cell != "" else None for cell in row] for row in rows]
    elif not none:
        rows = [[cell if cell is not None else "" for cell in row] for row in rows]
    return rows


def read_llist_from_xlsx_file(
    filename: str | Path | None = None,
    sheet: Any = None,
    cvt_dict: dict[Any, Any] | None = None,
    *,
    header: bool = False,
    none: bool = True,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as err:
        raise TableFormatError("openpyxl package is required for XLSX reading") from err

    workbook = None
    if hasattr(sheet, "iter_rows") and hasattr(sheet, "title"):
        worksheet = sheet
    else:
        if filename is None:
            raise TableFormatError("filename is required when sheet is not an XLSX worksheet object")
        workbook = load_workbook(filename, read_only=True, data_only=True)
        worksheet = workbook[sheet] if isinstance(sheet, str) else workbook.worksheets[sheet or 0]
    try:
        rows = [[cell for cell in row] for row in worksheet.iter_rows(values_only=True)]
    finally:
        if workbook is not None:
            workbook.close()

    rows = _normalize_spreadsheet_rows(
        rows,
        none=none,
        skip_empty=skip_empty,
        preserve_first=preserve_first,
        eol=eol,
        trim_empty_tails=trim_empty_tails,
    )
    _apply_table_conversion(rows, cvt_dict, header=header, source_name="XLSX")
    return rows


def _open_xls_workbook(xlrd_module: Any, filename: str | Path) -> Any:
    encoding_map = getattr(xlrd_module.book, "encoding_from_codepage", None)
    if isinstance(encoding_map, dict):
        encoding_map.setdefault(65001, "utf_8")
    try:
        return xlrd_module.open_workbook(filename)
    except UnicodeDecodeError as err:
        if "utf-8" not in str(err).lower():
            raise
        return xlrd_module.open_workbook(filename, encoding_override="cp1251")


def read_llist_from_xls_file(
    filename: str | Path | None = None,
    sheet: Any = None,
    cvt_dict: dict[Any, Any] | None = None,
    *,
    header: bool = False,
    none: bool = False,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    try:
        import xlrd
    except ImportError as err:
        raise TableFormatError("xlrd package is required for XLS reading") from err

    workbook = None
    if hasattr(sheet, "nrows") and hasattr(sheet, "row"):
        worksheet = sheet
    else:
        if filename is None:
            raise TableFormatError("filename is required when sheet is not an XLS worksheet object")
        workbook = _open_xls_workbook(xlrd, filename)
        worksheet = workbook.sheet_by_name(sheet) if isinstance(sheet, str) else workbook.sheet_by_index(sheet or 0)
    try:
        rows = [
            [worksheet.cell_value(row_index, column_index) for column_index in range(worksheet.ncols)]
            for row_index in range(worksheet.nrows)
        ]
    finally:
        if workbook is not None:
            workbook.release_resources()
    rows = _normalize_spreadsheet_rows(
        rows,
        none=none,
        skip_empty=skip_empty,
        preserve_first=preserve_first,
        eol=eol,
        trim_empty_tails=trim_empty_tails,
        empty_string_to_none=True,
    )
    _apply_table_conversion(rows, cvt_dict, header=header, source_name="XLS")
    return rows


def read_llist_from_file(
    filename: str | Path,
    sheet: str | int | None = None,
    cvt_dict: dict[Any, Any] | None = None,
    *,
    header: bool | None = False,
    none: bool | None = True,
    skip_empty: bool = True,
    preserve_first: bool = True,
    eol: Any = None,
    codepage: str = R_CODEPAGE,
    trim_empty_tails: bool = True,
) -> list[list[Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return read_llist_from_csv_file(
            filename,
            cvt_dict=cvt_dict,
            header=bool(header),
            codepage=codepage,
            none=bool(none),
            skip_empty=skip_empty,
            preserve_first=preserve_first,
            eol=eol,
            trim_empty_tails=trim_empty_tails,
        )
    if suffix == ".json":
        return read_llist_from_json_file(
            filename,
            cvt_dict=cvt_dict,
            header=bool(header),
            codepage=codepage,
            none=none,
            skip_empty=skip_empty,
            preserve_first=preserve_first,
            eol=eol,
            trim_empty_tails=trim_empty_tails,
        )
    if suffix == ".xml":
        return read_llist_from_xml_file(
            filename,
            cvt_dict=cvt_dict,
            header=bool(header),
            codepage=codepage,
            none=none,
            skip_empty=skip_empty,
            preserve_first=preserve_first,
            eol=eol,
            trim_empty_tails=trim_empty_tails,
        )
    if suffix == ".xlsx":
        return read_llist_from_xlsx_file(
            filename,
            sheet=sheet,
            cvt_dict=cvt_dict,
            header=bool(header),
            none=bool(none),
            skip_empty=skip_empty,
            preserve_first=preserve_first,
            eol=eol,
            trim_empty_tails=trim_empty_tails,
        )
    if suffix == ".xls":
        return read_llist_from_xls_file(
            filename,
            sheet=sheet,
            cvt_dict=cvt_dict,
            header=bool(header),
            none=bool(none),
            skip_empty=skip_empty,
            preserve_first=preserve_first,
            eol=eol,
            trim_empty_tails=trim_empty_tails,
        )
    raise TableFormatError(f"Unsupported table file extension: {suffix}")


def save_llist_to_csv_file(
    rows: list[list[Any]],
    filename: str | Path,
    *,
    codepage: str = UTF8_BOM,
) -> None:
    path = Path(add_ext(filename, ".csv"))
    with path.open("w", encoding=codepage, newline="") as file:
        writer = csv.writer(file)
        writer.writerows(rows)


def save_llist_to_txt_file(
    rows: list[list[Any]],
    filename: str | Path,
    *,
    codepage: str = W_CODEPAGE,
) -> None:
    path = Path(add_ext(filename, ".txt"))
    text = "\n".join("\t".join(str(item) for item in row) for row in rows)
    path.write_text(text, encoding=codepage)


def save_llist_to_xlsx_file(
    rows: list[list[Any]],
    filename: str | Path,
    *,
    sheet: str | None = None,
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as err:
        raise TableFormatError("openpyxl package is required for XLSX writing") from err
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise TableFormatError("XLSX workbook has no active worksheet")
    worksheet.title = sheet or "Sheet1"
    for row in rows:
        worksheet.append(row)
    workbook.save(add_ext(filename, ".xlsx"))


def save_llist_to_xls_file(
    rows: list[list[Any]],
    filename: str | Path,
    *,
    sheet: str | None = None,
) -> None:
    try:
        import xlwt  # type: ignore[import-untyped]
    except ImportError as err:
        raise TableFormatError("xlwt package is required for XLS writing") from err
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet(sheet or "Sheet1")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            worksheet.write(row_index, column_index, value)
    workbook.save(add_ext(filename, ".xls"))


def save_llist_to_file(
    rows: list[list[Any]],
    filename: str | Path,
    sheet: str | None = None,
    *,
    codepage: str = W_CODEPAGE,
) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        save_llist_to_csv_file(rows, filename, codepage=codepage)
        return
    if suffix == ".txt":
        save_llist_to_txt_file(rows, filename, codepage=codepage)
        return
    if suffix == ".xlsx":
        save_llist_to_xlsx_file(rows, filename, sheet=sheet)
        return
    if suffix == ".xls":
        save_llist_to_xls_file(rows, filename, sheet=sheet)
        return
    raise TableFormatError(f"Unsupported table file extension: {suffix}")


def is_opened(filename: str | Path) -> bool:
    path = Path(filename)
    if not path.exists():
        return False
    try:
        with path.open("a"):
            return False
    except OSError:
        return True
